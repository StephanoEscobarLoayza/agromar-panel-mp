"""Carga el historial de subproductos (emulsión y aceite) desde el Excel de la
planta ("SUBPRODUCTOS NARANJA CAMPAÑA-2026") a las tablas emulsion_* y aceite_*.

Uso:
    python importar_subproductos.py "ruta/al/excel.xlsx"             # solo revisa e imprime el informe
    python importar_subproductos.py "ruta/al/excel.xlsx" --aplicar   # además escribe en la base
    python importar_subproductos.py "ruta/al/excel.xlsx" --aplicar --reiniciar   # borra lo cargado antes y vuelve a cargar

Se lee la hoja "1. EMULSION-ACEITE NJ ORG" (cilindros de emulsión, fechas de
transformación y barriles de aceite) y la hoja "2. SEGUIMIENTO..." solo para la
ubicación de cada barril. Reglas:
  * cada bloque de la hoja (una fila "CORRIDA" con su MP procesada) es un
    registro de emulsión; se vincula a las corridas del panel cuya MP suma lo
    mismo (se prueban combinaciones de hasta 3 corridas de los 3 días previos);
  * la fecha de transformación de cada cilindro agrupa las transformaciones a
    aceite; el aceite de una transformación es la suma de los barriles
    anotados con esa fecha;
  * un barril puede aparecer en dos fechas ("SALDO PARA n"): son sus aportes; cada
    aporte se atribuye a la corrida (bloque) donde está anotado en la hoja.
Las fechas del Excel vienen mezcladas (texto y fechas con día y mes invertidos):
se interpreta la que cae más cerca de la fecha de referencia de la fila.
"""
import argparse
import datetime as dt
import itertools
import os
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from sqlalchemy import create_engine, text

BASE_DIR = Path(__file__).parent
load_dotenv(BASE_DIR / ".env")

ANIO = 2026                          # año de la campaña
AJUSTES_ANIO = []                    # celdas cuyo año venía mal escrito y se corrigió
CODIGO_BASE_BARRIL = 210475          # el barril n.º 1 es el 210476
TOLERANCIA_VINCULO = 0.06            # diferencia máxima (MP Excel vs suma de corridas) para vincular
UBICACIONES = {"PRECAMARA": "Precámara", "PRECÁMARA": "Precámara", "EXTRACCIÓN": "Extracción",
               "EXTRACCION": "Extracción", "REEFER 2": "Reefer 2"}


def fecha_celda(v, ref=None, despues=False):
    """Convierte una celda a date. Si es una fecha de Excel con día <= 12 puede
    estar invertida (día y mes): se elige la interpretación más cercana a ref
    (con despues=True se prefiere la que cae el mismo día o después de ref)."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        v = v.date()
    if isinstance(v, dt.date):
        cands = {v}
        if v.day <= 12:
            try:
                cands.add(dt.date(v.year, v.day, v.month))
            except ValueError:
                pass
        if ref is None or len(cands) == 1:
            return v
        return min(cands, key=lambda d: ((d < ref) if despues else False, abs((d - ref).days)))
    if isinstance(v, str):
        # puede venir un rango ("30/07/2026- 31/07/2026"): vale la última fecha
        todas = re.findall(r"(\d{1,2})/(\d{1,2})(?:/(\d{4}|\d{2}))?", v)
        if todas:
            d, mo, y = todas[-1]
            anio = int(y) if y else ANIO
            anio = anio + 2000 if anio < 100 else anio
            if anio != ANIO:                       # error de digitación del año (p. ej. 29/09/2027)
                AJUSTES_ANIO.append(v.strip())
                anio = ANIO
            return dt.date(anio, int(mo), int(d))
    return None


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def leer_hoja_principal(ws):
    """Lee la hoja de emulsión/aceite. Las fechas ambiguas se resuelven con la
    secuencia cronológica de la propia hoja (la fecha de emulsión avanza poco a
    poco; la de transformación va 0-5 días después de la de emulsión)."""
    bloques, cilindros, lineas_aceite, avisos = [], [], [], []
    actual = None
    corriente = None       # última fecha de emulsión aceptada
    for n_fila, f in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        f = list(f) + [None] * (20 - len(f))
        B, C, D, E, G, H, J, K, L, M, N, R = f[1], f[2], f[3], f[4], f[6], f[7], f[9], f[10], f[11], f[12], f[13], f[17]
        if B not in (None, "") and num(C) is not None:
            fecha = fecha_celda(B, corriente)
            actual = {"fila": n_fila, "fecha": fecha, "mp": num(C), "lotes": [], "cilindros": []}
            bloques.append(actual)
            if corriente is None:
                corriente = fecha
        if actual is None:
            continue
        if E not in (None, ""):
            actual["lotes"].append(str(int(E)) if isinstance(E, float) and E == int(E) else str(E).strip())
        f_emul = None
        if D not in (None, ""):
            f_emul = fecha_celda(D, corriente)
            if f_emul is not None:
                corriente = f_emul
        if isinstance(G, str) and G.upper().startswith("EORG"):
            numero = int(G.strip().split("-")[-1])
            f_emul = f_emul or corriente
            k = fecha_celda(K, f_emul, despues=True)
            descartada = None
            if k is not None and f_emul is not None and not (0 <= (k - f_emul).days <= 7):
                descartada, k = k, None     # una transformación no ocurre semanas después ni antes de la emulsión
            c = {"numero": numero, "fecha": f_emul, "peso": num(H), "transf": k, "k_descartada": descartada,
                 "estado_excel": (J or "").strip(), "bloque": len(bloques) - 1, "fila": n_fila}
            actual["cilindros"].append(c)
            cilindros.append(c)
        if M not in (None, "") and num(N) is not None:
            codigo = str(M).strip()
            if re.fullmatch(r"\d+", codigo):
                lineas_aceite.append({"codigo": codigo, "kg": num(N), "fecha": fecha_celda(K, f_emul or corriente, despues=True),
                                      "nisira": R if isinstance(R, bool) else None, "fila": n_fila, "bloque": len(bloques) - 1})
    # la fecha del bloque debe estar cerca de la fecha de emulsión de sus cilindros
    # (un error de digitación, p. ej. 09/09 en lugar de 09/08, se corrige con ella)
    for bl in bloques:
        fechas = [c["fecha"] for c in bl["cilindros"] if c["fecha"]]
        if not fechas:
            continue
        esperado = max(fechas)
        if bl["fecha"] is None or abs((bl["fecha"] - esperado).days) > 4:
            avisos.append(f"Bloque de la fila {bl['fila']}: la fecha «{bl['fecha']}» no concuerda con la de sus cilindros; se usó {esperado}.")
            bl["fecha"] = esperado
    # la fecha de emulsión de un cilindro no puede quedar semanas lejos de la de su bloque
    # (p. ej. «01/09» escrito en lugar de «01/10»): se toma la fecha del bloque
    for bl in bloques:
        corregidos = [c for c in bl["cilindros"] if bl["fecha"] and c["fecha"] and abs((c["fecha"] - bl["fecha"]).days) > 4]
        if corregidos:
            avisos.append(f"Cilindros {corregidos[0]['numero']} a {corregidos[-1]['numero']}: la fecha de emulsión anotada ({corregidos[0]['fecha']}) "
                          f"no concuerda con la de su corrida ({bl['fecha']}); se usó la de la corrida.")
            for c in corregidos:
                c["fecha"] = bl["fecha"]
    # un cilindro «Procesado» sin fecha de transformación (se olvidó anotarla): se le
    # asigna la del cilindro anterior (si es una fecha posible), y queda marcado
    previo_c = None
    for c in cilindros:
        if c["transf"] is None and c["estado_excel"].lower() == "procesado" and previo_c                 and previo_c["transf"] is not None and c["fecha"] is not None                 and 0 <= (previo_c["transf"] - c["fecha"]).days <= 7:
            c["transf"] = previo_c["transf"]
            c["inferido"] = True
        previo_c = c
    return bloques, cilindros, lineas_aceite, avisos


def leer_ubicaciones(wb):
    """codigo del barril -> ubicación, desde la hoja de seguimiento por cilindro."""
    ubic = {}
    for ws in wb.worksheets:
        if "SEGUIMIENTO" not in ws.title.upper():
            continue
        for f in ws.iter_rows(min_row=3, values_only=True):
            f = list(f) + [None] * (12 - len(f))
            codigo, u = f[8], f[11]
            if codigo not in (None, "") and u:
                ubic[str(codigo).strip()] = UBICACIONES.get(str(u).strip().upper(), str(u).strip().title())
    return ubic


MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
         "setiembre": 9, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}


def fecha_del_nombre(nombre):
    """«12-agosto JCC nj org» -> 12/08. La fecha de inicio guardada en una corrida puede
    estar mal escrita (p. ej. 12-agosto guardada con inicio el 08/08); el nombre es más fiable."""
    m = re.match(r"^\s*(\d{1,2})(?:\.\d+)?\s*[- ]\s*([A-Za-zñÑ]+)", nombre or "")
    if m and m[2].lower() in MESES:
        try:
            return dt.date(ANIO, MESES[m[2].lower()], int(m[1]))
        except ValueError:
            return None
    return None


def vincular_corridas(bloque, corridas, usadas):
    """Mejor combinación (1 a 3) de corridas de los 3 días previos cuya suma de kg se
    acerca a la MP del bloque. Devuelve (lista de corridas, suma) o ([], 0)."""
    if bloque["fecha"] is None or not bloque["mp"]:
        return [], 0.0
    ini, fin = bloque["fecha"] - dt.timedelta(days=3), bloque["fecha"]
    cand = [c for c in corridas if ini <= c["fecha"] <= fin and c["kg"] > 0 and c["id"] not in usadas]
    mejor, mejor_dif = None, None
    for k in (1, 2, 3):
        for combo in itertools.combinations(cand, k):
            suma = sum(c["kg"] for c in combo)
            dif = abs(suma - bloque["mp"])
            if mejor is None or dif < mejor_dif - 1e-6:
                mejor, mejor_dif = combo, dif
    if mejor is None or mejor_dif / bloque["mp"] > TOLERANCIA_VINCULO:
        # sin una suma que cuadre: se vinculan las corridas cuyo nombre lleva exactamente esa fecha
        # (la diferencia de kg queda a la vista en la pantalla para revisarla)
        mismas = [c for c in cand if fecha_del_nombre(c["nombre"]) == bloque["fecha"]]
        return mismas, sum(c["kg"] for c in mismas)
    return list(mejor), sum(c["kg"] for c in mejor)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--aplicar", action="store_true", help="escribe en la base (sin esto solo informa)")
    ap.add_argument("--reiniciar", action="store_true", help="borra lo cargado antes de volver a cargar")
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    wb = openpyxl.load_workbook(args.excel, data_only=True)
    hoja = next(ws for ws in wb.worksheets if ws.title.strip().startswith("1."))
    bloques, cilindros, lineas, avisos_lectura = leer_hoja_principal(hoja)
    ubic = leer_ubicaciones(wb)

    engine = create_engine(os.environ.get("DATABASE_URL_POOLED") or os.environ["DATABASE_URL"])
    with engine.connect() as c:
        corridas = [dict(r._mapping) for r in c.execute(text(
            "SELECT c.id, c.nombre, c.fecha_inicio::date AS fecha, COALESCE(SUM(a.kg_asignados), 0)::float AS kg "
            "FROM corridas c LEFT JOIN asignaciones a ON a.corrida_id = c.id GROUP BY c.id ORDER BY c.fecha_inicio"))]
        hay = c.execute(text("SELECT count(*) FROM emulsion_registros")).scalar()
    for co in corridas:      # la fecha del nombre manda sobre la fecha de inicio guardada
        co["fecha"] = fecha_del_nombre(co["nombre"]) or co["fecha"]
    if hay and not args.reiniciar:
        print(f"Ya hay {hay} registros de emulsión cargados. Use --reiniciar para volver a cargar desde cero.")
        return

    avisos = list(avisos_lectura)
    # ---- bloques -> corridas del panel ----
    usadas = set()
    print("\n=== BLOQUES (registros de emulsión) ===")
    print(f"{'fecha':<11}{'MP Excel':>11} {'cil':>4}  corridas vinculadas (suma)")
    for b in bloques:
        b["corridas"], b["suma"] = vincular_corridas(b, corridas, usadas)
        usadas.update(c["id"] for c in b["corridas"])
        if b["corridas"]:
            dif = b["suma"] - b["mp"]
            txt = " + ".join(c["nombre"] for c in b["corridas"]) + f"  ({b['suma']:,.0f}, dif {dif:+,.0f})"
            if abs(dif) > 5:
                avisos.append(f"Bloque {b['fecha']}: MP del Excel {b['mp']:,.0f} frente a {b['suma']:,.0f} de las corridas ({dif:+,.0f} kg).")
        else:
            txt = "— sin vincular"
            avisos.append(f"Bloque {b['fecha']} (MP {b['mp']:,.0f}): no se encontró una combinación de corridas que cuadre.")
        print(f"{b['fecha']!s:<11}{b['mp']:>11,.0f} {len(b['cilindros']):>4}  {txt}")

    # ---- transformaciones: cilindros y barriles agrupados por fecha ----
    trans = defaultdict(lambda: {"cilindros": [], "lineas": []})
    sin_fecha = [c for c in cilindros if c["transf"] is None]
    for c in cilindros:
        if c["transf"] is not None:
            trans[c["transf"]]["cilindros"].append(c)
    for ln in lineas:
        if ln["fecha"] is None:
            avisos.append(f"Fila {ln['fila']}: el barril {ln['codigo']} ({ln['kg']} kg) no tiene fecha de transformación legible.")
            continue
        trans[ln["fecha"]]["lineas"].append(ln)

    print("\n=== TRANSFORMACIONES A ACEITE ===")
    print(f"{'fecha':<11}{'cil':>4}{'emulsión kg':>13}{'aceite kg':>11}{'conv %':>8}")
    for f in sorted(trans):
        t = trans[f]
        em = sum(c["peso"] or 0 for c in t["cilindros"])
        ac = sum(l["kg"] for l in t["lineas"])
        conv = f"{ac / em * 100:6.1f}" if em and ac else "     —"
        marca = ""
        if not t["lineas"]:
            marca = "  <- cilindros sin aceite anotado"
            avisos.append(f"Transformación {f}: los cilindros {', '.join(str(c['numero']) for c in t['cilindros'])} tienen esa fecha pero no hay barriles de aceite anotados.")
        if not t["cilindros"]:
            marca = "  <- aceite sin cilindros"
            avisos.append(f"Transformación {f}: {ac:,.2f} kg de aceite pero ningún cilindro con esa fecha de transformación.")
        print(f"{f!s:<11}{len(t['cilindros']):>4}{em:>13,.1f}{ac:>11,.2f}{conv:>8}{marca}")

    # ---- barriles ----
    barriles = defaultdict(list)
    for ln in lineas:
        if ln["fecha"] is not None:
            barriles[ln["codigo"]].append(ln)
    print("\n=== BARRILES ===")
    incompletos = []
    for codigo in sorted(barriles, key=int):
        tot = sum(l["kg"] for l in barriles[codigo])
        if abs(tot - 183) > 0.01:
            incompletos.append((codigo, tot))
    print(f"{len(barriles)} barriles, del {min(barriles, key=int)} al {max(barriles, key=int)}; "
          f"{sum(1 for c in barriles if len(barriles[c]) > 1)} con aportes de dos transformaciones.")
    for codigo, tot in incompletos:
        detalle = ", ".join(f"{l['fecha']}: {l['kg']:g}" for l in barriles[codigo])
        print(f"  barril {codigo}: {tot:,.2f} kg de 183 ({detalle})" + ("  (abierto: es el último)" if codigo == max(barriles, key=int) else ""))
        if codigo != max(barriles, key=int):
            avisos.append(f"El barril {codigo} suma {tot:,.2f} kg y no 183.")

    print("\n=== CILINDROS ===")
    print(f"{len(cilindros)} cilindros; {sum(1 for c in cilindros if c['transf'])} con fecha de transformación; "
          f"{len(sin_fecha)} sin fecha (quedan en precámara).")
    raros = [c for c in sin_fecha if c["estado_excel"].lower() == "procesado"]
    for c in raros:
        avisos.append(f"Cilindro EORG-2026-{c['numero']} figura «Procesado» pero no tiene fecha de transformación (fila {c['fila']}).")
    for c in cilindros:
        if c.get("inferido"):
            motivo = (f"la fecha anotada ({c['k_descartada']}) no es posible" if c.get("k_descartada")
                      else "no tiene fecha de transformación")
            avisos.append(f"Cilindro EORG-2026-{c['numero']}: {motivo} en el Excel; se le asignó la del cilindro anterior ({c['transf']}).")
    numeros = sorted(c["numero"] for c in cilindros)
    faltan = sorted(set(range(numeros[0], numeros[-1] + 1)) - set(numeros))
    if faltan:
        avisos.append(f"Faltan en el Excel los cilindros: {faltan}.")
    if len(numeros) != len(set(numeros)):
        avisos.append("Hay códigos de cilindro repetidos en el Excel.")

    if AJUSTES_ANIO:
        avisos.append(f"{len(AJUSTES_ANIO)} fechas con el año mal escrito se tomaron como {ANIO}: {', '.join(sorted(set(AJUSTES_ANIO)))}.")
    print("\n=== AVISOS ===")
    for a in avisos:
        print(" -", a)
    if not avisos:
        print(" (ninguno)")

    if not args.aplicar:
        print("\nRevisión terminada: no se escribió nada. Use --aplicar para cargar.")
        return

    # ---------------------------------------------------------------- escritura
    with engine.begin() as c:
        if args.reiniciar:
            for t in ("aceite_barril_aportes", "emulsion_cilindros", "aceite_barriles", "aceite_transformaciones",
                      "emulsion_registro_corridas", "emulsion_registros"):
                c.execute(text(f"DELETE FROM {t}"))
        reg_id = {}
        for i, b in enumerate(bloques):
            obs = "Importado del Excel de subproductos." + (f" Lotes: {', '.join(b['lotes'])}." if b["lotes"] else "")
            reg_id[i] = c.execute(text(
                "INSERT INTO emulsion_registros (fecha, mp_kg_manual, observaciones) VALUES (:f, :mp, :o) RETURNING id"),
                {"f": b["fecha"], "mp": b["mp"], "o": obs}).scalar()
            for co in b["corridas"]:
                c.execute(text("INSERT INTO emulsion_registro_corridas (registro_id, corrida_id) VALUES (:r, :c)"),
                          {"r": reg_id[i], "c": co["id"]})
        trans_id = {}
        for f in sorted(trans):
            ac = sum(l["kg"] for l in trans[f]["lineas"])
            if ac <= 0:
                continue   # sin aceite anotado: sus cilindros quedan en precámara
            trans_id[f] = c.execute(text(
                "INSERT INTO aceite_transformaciones (fecha, aceite_kg, observaciones) VALUES (:f, :a, :o) RETURNING id"),
                {"f": f, "a": round(ac, 2), "o": "Importado del Excel de subproductos."}).scalar()
        for cil in cilindros:
            c.execute(text(
                "INSERT INTO emulsion_cilindros (numero, registro_id, fecha, peso_neto_kg, transformacion_id) "
                "VALUES (:n, :r, :f, :p, :t)"),
                {"n": cil["numero"], "r": reg_id[cil["bloque"]], "f": cil["fecha"], "p": cil["peso"],
                 "t": trans_id.get(cil["transf"])})
        for codigo in sorted(barriles, key=int):
            numero = int(codigo) - CODIGO_BASE_BARRIL
            nis = any(l["nisira"] for l in barriles[codigo])
            bid = c.execute(text(
                "INSERT INTO aceite_barriles (numero, codigo, ubicacion, nisira) VALUES (:n, :c, :u, :s) RETURNING id"),
                {"n": numero, "c": codigo, "u": ubic.get(codigo), "s": nis}).scalar()
            for l in sorted(barriles[codigo], key=lambda x: x["fecha"]):
                if l["fecha"] not in trans_id:
                    continue
                c.execute(text("INSERT INTO aceite_barril_aportes (barril_id, transformacion_id, registro_id, kg) VALUES (:b, :t, :r, :k)"),
                          {"b": bid, "t": trans_id[l["fecha"]], "r": reg_id[l["bloque"]], "k": round(l["kg"], 2)})
    print("\nCarga terminada.")


if __name__ == "__main__":
    main()
