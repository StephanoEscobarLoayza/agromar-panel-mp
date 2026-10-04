"""Excel de Subproductos (emulsión y aceite): resumen por corrida, cilindros de
emulsión y barriles de aceite. Mismo estilo visual que el Excel de Paradas
(título verde, acento cítrico, encabezados, cebreado y anchos pensados para leerse)."""
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_TITULO_FILL = PatternFill("solid", fgColor="1F4E37")
_ACENTO_FILL = PatternFill("solid", fgColor="E8890C")
_HEAD_FILL = PatternFill("solid", fgColor="CFE2C1")
_FILA_ALT_FILL = PatternFill("solid", fgColor="F3F8EF")
_AVISO_FILL = PatternFill("solid", fgColor="FBEBD6")
_HEAD_FONT = Font(bold=True, size=10, color="1F4E37")
_TITULO_FONT = Font(bold=True, size=13, color="FFFFFF")
_THIN = Side("thin", color="C9CEC5")
_BORDE = Border(_THIN, _THIN, _THIN, _THIN)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FMT_FECHA = "DD/MM/YYYY"
_FMT_KG = "#,##0.00"
_FMT_ENT = "#,##0"
_FMT_PCT = "0.0%"

_ESTADO_REG = {"PRECAMARA": "En precámara", "PARCIAL": "Parcial", "TRANSFORMADO": "Transformado"}


def _ahora_peru():
    return datetime.now(timezone.utc).astimezone(timezone(timedelta(hours=-5)))


def _hoja(wb, titulo, subtitulo, columnas, filas, primera=False):
    """columnas: [(nombre, ancho, formato|None, alineacion|None)]. filas: listas de valores
    (o (valores, resaltar) para pintar la fila en ámbar)."""
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo[:31]
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    n = len(columnas)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    c = ws.cell(row=1, column=1, value=titulo)
    c.font, c.fill = _TITULO_FONT, _TITULO_FILL
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    for i in range(1, n + 1):
        ws.cell(row=2, column=i).fill = _ACENTO_FILL
    ws.row_dimensions[2].height = 4
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n)
    c = ws.cell(row=3, column=1, value=subtitulo)
    c.font = Font(italic=True, size=10, color="5B6459")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 20

    h = 5
    for i, (nombre, ancho, _fmt, _al) in enumerate(columnas, start=1):
        cc = ws.cell(row=h, column=i, value=nombre)
        cc.font, cc.fill, cc.alignment, cc.border = _HEAD_FONT, _HEAD_FILL, _CENTRO, _BORDE
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[h].height = 30

    r = h + 1
    for idx, fila in enumerate(filas):
        valores, resaltar = fila if isinstance(fila, tuple) else (fila, False)
        relleno = _AVISO_FILL if resaltar else (_FILA_ALT_FILL if idx % 2 else None)
        for i, v in enumerate(valores, start=1):
            _n, _a, fmt, al = columnas[i - 1]
            cc = ws.cell(row=r, column=i, value=v)
            cc.border = _BORDE
            if relleno:
                cc.fill = relleno
            if fmt:
                cc.number_format = fmt
            cc.alignment = Alignment(horizontal=al or ("right" if fmt in (_FMT_KG, _FMT_ENT, _FMT_PCT) else "left"),
                                     vertical="top", wrap_text=True)
        r += 1
    if not filas:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
        ws.cell(row=r, column=1, value="Sin registros.").font = Font(italic=True, color="8B9186")
        r += 1
    ws.freeze_panes = f"A{h + 1}"
    ws.print_title_rows = f"{h}:{h}"
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
    c = ws.cell(row=r, column=1, value=f"Generado el {_ahora_peru().strftime('%d/%m/%Y %H:%M')} desde el Panel de cuadre de producción.")
    c.font = Font(italic=True, size=9, color="8B9186")
    c.alignment = Alignment(horizontal="left", indent=1)
    return ws


def generar_subproductos_xlsx(registros, cilindros, barriles, kpis) -> bytes:
    wb = Workbook()
    resumen = (f"{kpis['n_registros']} registros · {kpis['n_cilindros']} cilindros de emulsión ({kpis['emulsion_kg']:,.2f} kg) · "
               f"{kpis['aceite_kg']:,.2f} kg de aceite en {kpis['n_barriles']} barriles")

    # ---- 1. resumen por corrida ----
    cols = [("Fecha", 12, _FMT_FECHA, "left"), ("Corridas de origen", 34, None, None), ("MP procesada (kg)", 14, _FMT_KG, None),
            ("Suma de las corridas (kg)", 15, _FMT_KG, None), ("Diferencia (kg)", 13, _FMT_KG, None), ("Cilindros", 10, _FMT_ENT, None),
            ("Desde", 15, None, None), ("Hasta", 15, None, None), ("Emulsión (kg)", 13, _FMT_KG, None),
            ("Aceite (kg)", 13, _FMT_KG, None), ("Aceite / t de MP", 12, _FMT_KG, None), ("Conversión aceite / emulsión", 14, _FMT_PCT, None),
            ("Estado", 14, None, None), ("Observaciones", 36, None, None)]
    filas = []
    for r in registros:
        dif = r["diferencia_kg"]
        filas.append(([
            r["fecha"], ", ".join(c["nombre"] for c in r["corridas"]) or "Sin vincular",
            r["mp_kg"], r["mp_corridas_kg"], dif, r["n_cilindros"],
            f"EORG-2026-{r['primer_cilindro']}" if r["primer_cilindro"] else None,
            f"EORG-2026-{r['ultimo_cilindro']}" if r["ultimo_cilindro"] else None,
            r["emulsion_kg"], r["aceite_kg"], r["aceite_por_t_mp"], r["conversion"],
            _ESTADO_REG.get(r["estado"], r["estado"]), r["observaciones"],
        ], dif is not None and abs(dif) > 5))
    _hoja(wb, "Resumen por corrida", resumen, cols, filas, primera=True)

    # ---- 2. cilindros de emulsión ----
    cols = [("Código", 16, None, None), ("Fecha de emulsión", 14, _FMT_FECHA, "left"), ("Corridas de origen", 34, None, None),
            ("Peso neto (kg)", 13, _FMT_KG, None), ("Estado", 13, None, None), ("Transformado a aceite", 15, _FMT_FECHA, "left")]
    filas = [[c["codigo"], c["fecha"], ", ".join(c["corridas"]) or "Sin vincular", c["peso_neto_kg"],
              "En precámara" if c["estado"] == "PRECAMARA" else "Procesado", c["transformacion_fecha"]]
             for c in cilindros]
    _hoja(wb, "Cilindros de emulsión", resumen, cols, filas)

    # ---- 3. barriles de aceite ----
    cols = [("N.º", 7, _FMT_ENT, "left"), ("Código", 12, None, None), ("Aceite (kg)", 12, _FMT_KG, None),
            ("Capacidad (kg)", 12, _FMT_KG, None), ("Origen del aceite (corrida · kg)", 30, None, None),
            ("Corridas del panel", 40, None, None), ("Trazabilidad", 20, None, None), ("Ubicación", 13, None, None),
            ("Registrado en Nisira", 12, None, "center")]
    filas = []
    for b in barriles:
        origen = "\n".join(f"{(a['registro_fecha'] or a['fecha']).strftime('%d/%m/%Y')} · {a['kg']:,.2f} kg" for a in b["aportes"])
        filas.append(([b["numero"], b["codigo"], b["kg"], b["capacidad_kg"], origen, ", ".join(b["corridas"]) or "—",
                       f"Mezcla de {b['n_origenes']} corridas" if b["mezcla"] else "Única", b["ubicacion"] or "—",
                       "Sí" if b["nisira"] else "No"],
                      b["kg"] > b["capacidad_kg"] + 0.005))
    _hoja(wb, "Barriles de aceite", resumen, cols, filas)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
