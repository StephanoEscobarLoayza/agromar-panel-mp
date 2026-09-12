"""Genera el .xlsx de "Paradas por período" (botón en Paradas). Antes se
armaba con el mismo helper genérico del export de respaldo (`_hoja_xlsx`
en api.py) - una hoja técnica, con los nombres de columna tal cual la base
de datos y sin ajustar anchos al contenido, que Stephano vio "poco
profesional" (fechas mostrando `#######`, columnas de texto largo cortadas).
Esta versión sí está pensada para leerse/imprimirse: título con el período,
resumen arriba, encabezados en español y columnas con ancho + wrap
pensados para lo que de verdad se escribe en cada campo."""
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_TITULO_FILL = PatternFill("solid", fgColor="1F4E37")   # verde bosque de la marca
_ACENTO_FILL = PatternFill("solid", fgColor="E8890C")   # cítrico de la marca - acento bajo el título, mismo criterio que las secciones de los PDF
_HEAD_FILL = PatternFill("solid", fgColor="CFE2C1")
_FILA_ALT_FILL = PatternFill("solid", fgColor="F3F8EF")  # cebreado sutil de filas, para que no se vea una pared plana en períodos largos
_EN_CURSO_FILL = PatternFill("solid", fgColor="FBEBD6")  # ámbar - misma paleta "warn" que ya usa la app, para no perder de vista una parada sin cerrar
_BOLD = Font(bold=True)
_HEAD_FONT = Font(bold=True, size=10, color="1F4E37")
_TITULO_FONT = Font(bold=True, size=13, color="FFFFFF")
_THIN = Side("thin", color="C9CEC5")
_BORDE = Border(_THIN, _THIN, _THIN, _THIN)
_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FMT_FECHAHORA = "DD/MM/YYYY HH:MM"

COLS = [
    ("Corrida", 20),
    ("Turno", 9),
    ("Inicio", 17),
    ("Fin", 17),
    ("Minutos", 10),
    ("Área / proceso", 16),
    ("Tipo de parada", 15),
    ("Equipo afectado", 18),
    ("Descripción de la falla", 34),
    ("Responsable", 14),
    ("Solución / observaciones", 34),
    ("Recomendación", 30),
]


def _ahora_peru():
    return datetime.now(timezone.utc).astimezone(timezone(timedelta(hours=-5)))


def _fmt_fecha_corta(d):
    return d.strftime("%d/%m/%Y")


def _fmt_minutos(min_total):
    if min_total is None:
        return "en curso"
    h, m = divmod(int(round(min_total)), 60)
    return f"{h}h {m:02d}m" if h else f"{m}m"


def generar_paradas_periodo_xlsx(desde, hasta, paradas) -> bytes:
    """desde/hasta: date. paradas: filas (corrida, turno, hora_inicio,
    hora_fin, min_total, area_proceso, tipo_parada, equipo_afectado,
    descripcion_falla, responsable_solucion, solucion_obs, recomendacion)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Paradas"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ncols = len(COLS)

    # ---------- título ----------
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=f"Paradas del {_fmt_fecha_corta(desde)} al {_fmt_fecha_corta(hasta)}")
    c.font = _TITULO_FONT
    c.fill = _TITULO_FILL
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 26

    # acento cítrico bajo el título - mismo lenguaje visual que la doble
    # regla forest+citrus de los encabezados de sección en los PDF
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    for i in range(1, ncols + 1):
        ws.cell(row=2, column=i).fill = _ACENTO_FILL
    ws.row_dimensions[2].height = 4

    # ---------- resumen ----------
    total_min = sum(float(p.get("min_total") or 0) for p in paradas)
    n = len(paradas)
    en_curso = sum(1 for p in paradas if p.get("hora_fin") is None)
    resumen = f"{n} parada{'s' if n != 1 else ''} · {_fmt_minutos(total_min)} parados en total"
    if en_curso:
        resumen += f" · {en_curso} en curso"
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    c = ws.cell(row=3, column=1, value=resumen)
    c.font = Font(italic=True, size=10, color="5B6459")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 20

    # ---------- encabezados ----------
    hfila = 5
    for i, (nombre, ancho) in enumerate(COLS, start=1):
        cc = ws.cell(row=hfila, column=i, value=nombre)
        cc.font = _HEAD_FONT
        cc.fill = _HEAD_FILL
        cc.alignment = _CENTRO
        cc.border = _BORDE
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[hfila].height = 28

    # ---------- filas ----------
    r = hfila + 1
    if not paradas:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        cc = ws.cell(row=r, column=1, value="Sin paradas registradas en este período.")
        cc.font = Font(italic=True, color="8B9186")
        cc.alignment = Alignment(horizontal="center")
    for idx, p in enumerate(paradas):
        fin = p.get("hora_fin")
        # una parada sin cerrar se resalta entera en ámbar (misma paleta
        # "warn" que ya usa el resto de la app) - no hay que leer cada
        # fila para notar que todavía sigue abierta. Si no, cebreado
        # sutil cada 2 filas para que un período largo no se vea plano.
        fila_fill = _EN_CURSO_FILL if fin is None else (_FILA_ALT_FILL if idx % 2 == 1 else None)
        fila = [
            p.get("corrida") or "—",
            "Día" if p.get("turno") == "DÍA" else "Noche" if p.get("turno") == "NOCHE" else (p.get("turno") or "—"),
            p.get("hora_inicio"),
            fin,
            p.get("min_total"),
            p.get("area_proceso") or "—",
            p.get("tipo_parada") or "—",
            p.get("equipo_afectado") or "—",
            p.get("descripcion_falla") or "—",
            p.get("responsable_solucion") or "—",
            p.get("solucion_obs") or "—",
            p.get("recomendacion") or "—",
        ]
        for i, v in enumerate(fila, start=1):
            cc = ws.cell(row=r, column=i, value=v)
            cc.border = _BORDE
            if fila_fill:
                cc.fill = fila_fill
            if i in (3, 4):
                cc.number_format = _FMT_FECHAHORA
                cc.alignment = Alignment(vertical="top")
            elif i == 5:
                cc.alignment = Alignment(horizontal="right", vertical="top")
                if v is None:
                    cc.value = "en curso"
                    cc.font = _BOLD
                    cc.alignment = Alignment(horizontal="center", vertical="top")
            elif i in (9, 11, 12):
                cc.alignment = _WRAP
            else:
                cc.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1

    ws.freeze_panes = f"A{hfila + 1}"
    ws.print_title_rows = f"{hfila}:{hfila}"  # repite el encabezado en cada página al imprimir

    # ---------- pie ----------
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    cc = ws.cell(row=r, column=1, value=f"Generado el {_ahora_peru().strftime('%d/%m/%Y %H:%M')} desde el Panel de cuadre de producción.")
    cc.font = Font(italic=True, size=9, color="8B9186")
    cc.alignment = Alignment(horizontal="left", indent=1)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
