"""Pruebas de las piezas que no dependen de la base de datos.

Se ejecutan con:  python -m unittest discover -s tests
No se abre ninguna conexión: el motor de SQLAlchemy se crea sin conectarse.
"""
import os
import re
import sys
import time
import logging
import unittest
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
logging.disable(logging.CRITICAL)   # los errores simulados no deben ensuciar la salida
sys.path.insert(0, str(RAIZ))
os.environ.setdefault("DATABASE_URL", "postgresql://usuario:clave@localhost:5432/prueba")
os.environ["APP_USER"] = os.environ["APP_PASSWORD"] = os.environ["APP_SESSION_SECRET"] = "prueba"

import api  # noqa: E402
import repartir  # noqa: E402
from fastapi import HTTPException  # noqa: E402
from reporte_trazabilidad import _dia_mes_del_nombre  # noqa: E402


class _ErrorBD(Exception):
    """Imita el error que entrega el controlador de Postgres (pg8000 lo entrega como diccionario)."""

    def __init__(self, codigo, mensaje):
        super().__init__({"C": codigo, "M": mensaje})


class _Envuelto(Exception):
    """Imita el error de SQLAlchemy, que guarda el original en .orig."""

    def __init__(self, orig):
        super().__init__(str(orig))
        self.orig = orig


class ErroresDeBaseDeDatos(unittest.TestCase):
    def test_regla_escrita_en_la_base_se_muestra_tal_cual(self):
        e = api._error_http(_Envuelto(_ErrorBD("P0001", "El lote 2730 solo tiene 2220 kg de saldo.")))
        self.assertEqual((e.status_code, e.detail), (400, "El lote 2730 solo tiene 2220 kg de saldo."))

    def test_duplicado_da_mensaje_claro(self):
        e = api._error_http(_Envuelto(_ErrorBD("23505", "duplicate key value violates unique constraint")))
        self.assertEqual(e.status_code, 409)
        self.assertNotIn("constraint", e.detail)

    def test_error_inesperado_no_filtra_detalles_internos(self):
        e = api._error_http(_Envuelto(_ErrorBD("XX000", "relation \"secreta\" does not exist")))
        self.assertEqual(e.status_code, 500)
        self.assertNotIn("secreta", e.detail)


class Sesion(unittest.TestCase):
    def test_token_recien_emitido_es_valido(self):
        self.assertTrue(api._sesion_valida(api._token_sesion()))

    def test_token_alterado_no_vale(self):
        exp, azar, firma = api._token_sesion().split(".")
        self.assertFalse(api._sesion_valida(f"{int(exp) + 3600}.{azar}.{firma}"))
        self.assertFalse(api._sesion_valida(f"{exp}.{azar}.{'0' * len(firma)}"))

    def test_token_vencido_no_vale(self):
        exp = int(time.time()) - 10
        self.assertFalse(api._sesion_valida(f"{exp}.abc.{api._firma(exp, 'abc')}"))

    def test_basura_no_vale(self):
        for token in ("", "sin-punto", "abc.def", "a.b.c", None):
            self.assertFalse(api._sesion_valida(token), token)

    def test_cerrar_sesion_invalida_solo_ese_token(self):
        uno, otro = api._token_sesion(), api._token_sesion()
        api._sesiones_cerradas.add(uno.rsplit(".", 1)[1])
        self.assertFalse(api._sesion_valida(uno))
        self.assertTrue(api._sesion_valida(otro))      # la sesión de otra persona sigue abierta


class FechaEnElNombreDeLaCorrida(unittest.TestCase):
    def test_nombres_validos(self):
        self.assertEqual(_dia_mes_del_nombre("17-Octubre JCC"), (17, 10))
        self.assertEqual(_dia_mes_del_nombre("19-Octubre JCC TASTE"), (19, 10))
        self.assertEqual(_dia_mes_del_nombre("5 septiembre ICE"), (5, 9))

    def test_nombres_sin_fecha(self):
        for nombre in ("Corrida de prueba", "32-Octubre", "", None):
            self.assertIsNone(_dia_mes_del_nombre(nombre), nombre)


class ExcelConVariasCorridas(unittest.TestCase):
    def _hoja(self, nombre, tipo):
        from datetime import datetime
        corrida = {"nombre": nombre, "tipo_proceso": tipo, "fecha_inicio": datetime(2026, 10, 17, 8, 0), "fecha_final": None, "saldo_por_secuencia": True}
        return {"corrida": corrida, "lotes": [], "productos": [], "mediciones": [], "stock": []}

    def test_nombres_de_hoja_validos_y_unicos(self):
        from reporte_trazabilidad import _nombres_de_hoja
        res = _nombres_de_hoja(["17-Octubre JCC", "17-Octubre JCC", "A/B:C*D?[E]", "", "X" * 50])
        self.assertEqual(res[0], "17-Octubre JCC")
        self.assertEqual(res[1], "17-Octubre JCC (2)")
        self.assertNotRegex(res[2], r"[\[\]:*?/\\]")
        self.assertEqual(res[3], "Corrida")
        self.assertTrue(all(len(n) <= 31 for n in res))
        self.assertEqual(len({n.lower() for n in res}), len(res))

    def test_solo_la_ultima_hoja_lleva_el_cuadro_de_stock(self):
        import io
        import openpyxl
        from reporte_trazabilidad import generar_trazabilidad_grupo_xlsx
        libro = generar_trazabilidad_grupo_xlsx([self._hoja("16-Octubre ICEGEN", "ICEGEN"), self._hoja("17-Octubre ICEGEN", "ICEGEN"), self._hoja("17-Octubre JCC", "JCC")])
        wb = openpyxl.load_workbook(io.BytesIO(libro))
        self.assertEqual(wb.sheetnames, ["16-Octubre ICEGEN", "17-Octubre ICEGEN", "17-Octubre JCC"])
        textos = [{str(c.value) for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)} for ws in wb.worksheets]
        con_cuadro = [any(t.startswith("Stock inicial") for t in hoja) for hoja in textos]
        self.assertEqual(con_cuadro, [False, False, True])
        self.assertTrue(all(any("está en la hoja «17-Octubre JCC»" in t for t in hoja) for hoja in textos[:2]))


class TanqueOlvidado(unittest.TestCase):
    def test_numero_de_tanque(self):
        self.assertEqual(api._numero_tanque("TK12"), 12)
        self.assertEqual(api._numero_tanque(" tk 07 "), 7)
        for nombre in ("Tanque 1", "TK12-B", "", None):
            self.assertIsNone(api._numero_tanque(nombre), nombre)

    def test_queda_entre_el_anterior_y_el_siguiente(self):
        from datetime import datetime, timezone
        h = lambda hh, mm: datetime(2026, 10, 24, hh, mm, tzinfo=timezone.utc)
        vecinos = [("TK10", h(2, 0)), ("TK11", h(2, 10)), ("TK13", h(2, 40)), ("TK14", h(3, 0))]
        hora = api._instante_para_tanque(12, vecinos)
        self.assertTrue(h(2, 10) < hora < h(2, 40))

    def test_el_ultimo_tanque_usa_la_hora_de_ahora(self):
        from datetime import datetime, timezone
        vecinos = [("TK1", datetime(2026, 10, 24, 2, 0, tzinfo=timezone.utc))]
        self.assertIsNone(api._instante_para_tanque(2, vecinos))
        self.assertIsNone(api._instante_para_tanque(1, []))

    def test_el_primero_que_falta_queda_antes_del_siguiente(self):
        from datetime import datetime, timezone
        h = datetime(2026, 10, 24, 2, 0, tzinfo=timezone.utc)
        self.assertTrue(api._instante_para_tanque(1, [("TK2", h)]) < h)


class OrdenDeUnReparto(unittest.TestCase):
    def test_la_jsa_jcc_queda_al_final_en_los_mismos_lugares(self):
        seq = [160, 171, 172, 173, 174]                      # 171 ICEGEN, 172 JCC, 173 ICEGEN
        tipos = {171: "ICEGEN", 172: "JCC", 173: "ICEGEN"}
        orden, nueva = api._secuencia_con_principal_al_final(seq, [172, 171, 173], tipos)
        self.assertEqual(orden, [171, 173, 172])
        self.assertEqual(nueva, [160, 171, 173, 172, 174])

    def test_si_ya_estaba_ordenado_no_cambia(self):
        seq = [1, 2, 3]
        orden, nueva = api._secuencia_con_principal_al_final(seq, [2, 3], {2: "ICEGEN", 3: "JSA"})
        self.assertEqual((orden, nueva), ([2, 3], seq))

    def test_con_dos_principales_se_respeta_el_orden_de_la_lista(self):
        seq = [10, 11, 12]
        orden, _ = api._secuencia_con_principal_al_final(seq, [12, 11], {11: "JSA", 12: "JCC"})
        self.assertEqual(orden, [11, 12])


class ListaDeCorridasParaRepartir(unittest.TestCase):
    def test_ids_validos(self):
        self.assertEqual(repartir._ids("3,1,3"), [1, 3])

    def test_ids_invalidos(self):
        for cadena in ("1", "a,b", "1,2,3,4,5", ""):
            with self.assertRaises(HTTPException, msg=cadena):
                repartir._ids(cadena)


class PaginasWeb(unittest.TestCase):
    """Reglas de presentación que no deben volver a romperse."""

    paginas = sorted((RAIZ / "web").glob("*.html"))

    def test_no_se_usan_cuadros_nativos_del_navegador(self):
        for p in self.paginas:
            self.assertIsNone(re.search(r"(?<![\w.])alert\(", p.read_text(encoding="utf-8")), p.name)

    def test_el_zoom_no_esta_bloqueado(self):
        for p in self.paginas:
            self.assertNotIn("maximum-scale", p.read_text(encoding="utf-8"), p.name)

    def test_las_paginas_con_avisos_cargan_app_js(self):
        for p in self.paginas:
            texto = p.read_text(encoding="utf-8")
            if "avisar(" in texto:
                self.assertIn("app.js", texto, p.name)


if __name__ == "__main__":
    unittest.main()
